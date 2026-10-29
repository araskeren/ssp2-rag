import io

import pytest
from pypdf import PdfWriter

from ai_service.rag.chunking import chunk_document
from ai_service.rag.parsers import (
    PageText,
    ParsedDocument,
    ParseError,
    normalize_text,
    parse_bytes,
    strip_repeated_lines,
)


def test_parse_markdown_title_and_text():
    parsed = parse_bytes("panduan.md", b"# Judul Panduan\n\nIsi   paragraf.\r\n\r\n\r\nLagi.")
    assert parsed.title == "Judul Panduan"
    assert parsed.pages[0].page is None
    assert "Isi paragraf." in parsed.pages[0].text
    assert "\n\n\n" not in parsed.pages[0].text


def test_parse_txt_encoding_fallback():
    parsed = parse_bytes("catatan.txt", "Kasus déjà vu".encode("cp1252"))
    assert "déjà" in parsed.full_text


def test_unsupported_and_empty():
    with pytest.raises(ParseError, match="tidak didukung"):
        parse_bytes("data.xlsx", b"x")
    with pytest.raises(ParseError, match="kosong"):
        parse_bytes("kosong.txt", b"   ")


def test_scanned_pdf_detected():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    with pytest.raises(ParseError, match="pindaian"):
        parse_bytes("scan.pdf", buf.getvalue())


def test_docx_headings_and_tables():
    import docx

    document = docx.Document()
    document.add_heading("Bab Satu", level=1)
    document.add_paragraph("Paragraf pertama.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Peran"
    table.rows[0].cells[1].text = "Akses"
    buf = io.BytesIO()
    document.save(buf)
    parsed = parse_bytes("panduan.docx", buf.getvalue())
    text = parsed.full_text
    assert "# Bab Satu" in text and "Peran\tAkses" in text


def test_strip_repeated_header_footer():
    pages = [
        PageText(i, f"Kementerian Kesehatan RI\nIsi halaman {i}\nHalaman {i}") for i in range(1, 6)
    ]
    cleaned = strip_repeated_lines(pages)
    assert all("Kementerian Kesehatan RI" not in p.text for p in cleaned)
    assert all(f"Isi halaman {p.page}" in p.text for p in cleaned)


def test_full_text_roundtrip_preserves_pages():
    parsed = ParsedDocument(
        pages=[PageText(1, "hal satu"), PageText(2, "hal dua")],
        mime_type="application/pdf",
        page_count=2,
    )
    restored = ParsedDocument.from_full_text(parsed.full_text, "application/pdf")
    assert [p.page for p in restored.pages] == [1, 2]


def test_chunking_sections_and_sizes():
    body = (
        "# Bab A\n\n"
        + ("Kalimat tentang pelaporan kasus. " * 40)
        + "\n\n## Bab B\n\n"
        + ("Kalimat tentang spesimen. " * 40)
    )
    parsed = parse_bytes("doc.md", body.encode())
    chunks = chunk_document(parsed, chunk_size=300, chunk_overlap=50)
    assert len(chunks) >= 4
    assert all(len(c.content) <= 300 + 40 for c in chunks)  # + prefiks judul bagian
    assert chunks[0].section == "Bab A"
    assert any(c.section == "Bab B" and c.content.startswith("Bab B") for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    assert normalize_text("a \n\n\n\nb") == "a\n\nb"


def test_build_or_query_drops_stopwords_and_dupes():
    from ai_service.rag.store import build_or_query

    assert (
        build_or_query("Bagaimana alur pengiriman spesimen ke laboratorium? spesimen!")
        == "alur | pengiriman | spesimen | laboratorium"
    )
    assert build_or_query("di ke dari") == "__kosong__"
    assert "'" not in build_or_query("it's a 'test'")
